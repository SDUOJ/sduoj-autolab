import hashlib
import json
import os
from datetime import datetime, timedelta
from decimal import Decimal

import pytest
from cryptography.fernet import Fernet
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.dialects.mysql import BIGINT, LONGTEXT, TINYINT
from sqlalchemy.ext.compiler import compiles
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from db import AutoTaskRun, AutoTaskRunLog, Base, ProblemGroup, ProblemObjective, ProblemSet, ProblemSetAnswerSheet, ProblemSetAnswerSheetDetail, ProblemSubjective, ojCourse
from auto_task import store
from auto_task.configuration import ProviderConfig, provider, save_provider
from auto_task.entities import AiFile, AiFileLink, AiLedger, AiResource, AiReview, AiRun, AiSetting, TABLES, TaskJob
from auto_task.handlers import apply_result, plan_exam
from auto_task.resources import canonical_question
from auto_task.worker import run_once
from server import auto_task_ai as api


@compiles(LONGTEXT, "sqlite")
def sqlite_text(*_, **kwargs):
    return "TEXT"


@compiles(BIGINT, "sqlite")
@compiles(TINYINT, "sqlite")
def sqlite_integer(*_, **kwargs):
    return "INTEGER"


USER = {"userId": 42, "username": "student", "roles": ["user"], "groups": []}
ADMIN = {"userId": 7, "username": "teacher", "roles": ["admin"], "groups": []}


@pytest.fixture
def database(monkeypatch, tmp_path):
    test_url = os.environ.get("AUTOLAB_AI_TEST_DATABASE_URL")
    if test_url:
        from sqlalchemy.engine.url import make_url
        url = make_url(test_url)
        assert url.host in ("127.0.0.1", "localhost") and url.database == "autolab_ai_test"
        engine = create_engine(test_url)
    else:
        engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    tables = [AutoTaskRun.__table__, AutoTaskRunLog.__table__, ProblemGroup.__table__, ProblemObjective.__table__, ProblemSet.__table__,
              ProblemSetAnswerSheet.__table__, ProblemSetAnswerSheetDetail.__table__, ProblemSubjective.__table__, ojCourse.__table__, *TABLES]
    Base.metadata.create_all(engine, tables=tables)
    monkeypatch.setattr(store, "_factory", sessionmaker(bind=engine, expire_on_commit=False))
    monkeypatch.setenv("AUTOLAB_AI_DATA_DIR", str(tmp_path))
    monkeypatch.setenv("AUTOLAB_AI_SECRET_KEY", Fernet.generate_key().decode())
    monkeypatch.setenv("AUTOLAB_AI_WAKE", "0")
    monkeypatch.setenv("AUTOLAB_AI_VALIDATE_IDENTITY", "0")
    from auto_task import access
    actors = {"7": ADMIN, "42": USER, "8": {"userId": 8, "username": "collaborator", "roles": ["user"], "groups": [10]}}
    monkeypatch.setattr(access, "fetch_actor", lambda subject, fresh=False: dict(actors[str(subject)]))
    with store.transaction() as s:
        s.add(AiSetting(id="runtime", data='{"max_concurrency":2,"per_user_concurrency":1}'))
    yield engine
    Base.metadata.drop_all(engine, tables=tables)
    engine.dispose()


@pytest.fixture
def client(database):
    app = FastAPI()
    app.include_router(api.router)
    from server.auto_task_actions import router as actions
    app.include_router(actions, prefix="/auto-task/ai")
    from server.auto_task_native import router as native
    app.include_router(native, prefix="/auto-task/ai")
    from server.auto_task_questions import router as question_workflows
    app.include_router(question_workflows, prefix="/auto-task/ai")
    app.dependency_overrides[api.cover_header] = lambda: USER
    return TestClient(app)


def resource(s, owner="42", kind="collection", data=None, owner_username="student"):
    row = AiResource(id="res-" + owner, kind=kind, owner_id=owner, owner_username=owner_username, data=store.dumps(data or {"text": "2+2?"}))
    s.add(row)
    s.flush()
    return row


def enqueue(key="op", data=None):
    with store.transaction() as s:
        row = s.query(AiResource).first() or resource(s)
        job = store.submit(s, USER, "question.extract", data or {"resource_version": row.version}, key, row)
        s.flush()
        return job.id


def test_atomic_submission_and_legacy_projection(database):
    job_id = enqueue()
    with store.transaction() as s:
        assert s.query(AutoTaskRun).filter_by(id=job_id).one().status == "pending"
        assert s.query(AutoTaskRunLog).filter_by(task_id=job_id, tag="payload").count() == 1
        assert s.query(TaskJob).count() == 1


def test_document_reimport_by_content_reuses_existing_resource(client):
    body = {"kind": "document", "data": {"title": "Lecture", "text": "same lecture content"}}
    first = client.post("/auto-task/ai/resources/create", json=body).json()["data"]
    repeated = client.post("/auto-task/ai/resources/create", json=body).json()["data"]
    assert repeated["replay"] and repeated["id"] == first["id"]
    changed = client.post("/auto-task/ai/resources/create", json={**body, "data": {"title": "Lecture", "text": "different"}}).json()["data"]
    assert changed["id"] != first["id"]
    client.post("/auto-task/ai/resources/" + first["id"] + "/delete")
    assert client.post("/auto-task/ai/resources/create", json=body).status_code == 409


def test_same_key_replays_and_different_input_conflicts(database):
    assert enqueue() == enqueue()
    with pytest.raises(HTTPException) as error:
        enqueue(data={"different": True})
    assert error.value.status_code == 409


def test_rolled_back_submission_is_not_claimable(database):
    with pytest.raises(RuntimeError):
        with store.transaction() as s:
            store.submit(s, USER, "question.extract", {}, "rollback", resource(s))
            raise RuntimeError()
    assert store.claim() is None


def test_claim_fences_old_worker_after_lease_recovery(database):
    job_id = enqueue()
    first = store.claim()
    with store.transaction() as s:
        s.query(TaskJob).filter_by(id=job_id).update({"lease_until": datetime.utcnow() - timedelta(seconds=1)})
    second = store.claim()
    assert second["attempt"] == 2
    assert first["lease_token"] != second["lease_token"]
    with pytest.raises(store.LeaseLost):
        store.complete(job_id, first["lease_token"], {})
    store.complete(job_id, second["lease_token"], {"ok": True})


def test_resource_and_user_serialization(database):
    enqueue("one")
    enqueue("two")
    first = store.claim()
    assert store.claim() is None
    store.complete(first["id"], first["lease_token"], {})
    assert store.claim() is not None


def test_cancel_blocks_stale_side_effect(database):
    job_id = enqueue()
    active = store.claim()
    with store.transaction() as s:
        store.cancel(s, s.query(TaskJob).get(job_id), "cancel")
    with pytest.raises(store.LeaseLost):
        store.complete(job_id, active["lease_token"], {})
    assert store.claim() is None


def test_failed_job_backoff(database):
    job_id = enqueue()
    active = store.claim()
    store.fail(job_id, active["lease_token"], "temporary")
    assert store.claim() is None
    assert store.claim(datetime.utcnow() + timedelta(minutes=1))["attempt"] == 2


def test_stale_resource_result_is_rejected(database):
    job_id = enqueue()
    active = store.claim()
    with store.transaction() as s:
        s.query(AiResource).first().version += 1
    with pytest.raises(ValueError):
        store.complete(job_id, active["lease_token"], {}, lambda s, j, r: apply_result(s, j, r, 1))
    with store.transaction() as s:
        assert s.query(TaskJob).get(job_id).status == "running"


def test_points_exactly_once(database):
    with store.transaction() as s:
        store.ledger(s, "42", "grant", 20, "7", "grant")
    with store.transaction() as s:
        store.ledger(s, "42", "grant", 20, "7", "grant")
        assert store.account(s, "42").balance == Decimal(20)
        assert s.query(AiLedger).count() == 1


def test_negative_points_dont_suspend_account(database):
    with store.transaction() as s:
        store.ledger(s, "42", "usage", -1, "42", "usage")
    with pytest.raises(HTTPException) as error:
        enqueue()
    assert error.value.status_code == 402
    with store.transaction() as s:
        assert store.account(s, "42").status == "active"


def test_three_warnings_cancel_tasks_and_require_review(database):
    job_id = enqueue()
    active = store.claim()
    with store.transaction() as s:
        for index in range(3):
            s.add(AiReview(id=str(index), subject_id="42", decision="warning", reason="reason"))
        s.flush()
        assert store.recalculate_access(s, "42").status == "suspended"
    with pytest.raises(store.LeaseLost):
        store.heartbeat(job_id, active["lease_token"])
    with store.transaction() as s:
        s.query(AiReview).get("0").decision = "normal"
        s.flush()
        assert store.recalculate_access(s, "42").status == "suspended"
        assert store.recalculate_access(s, "42", True).status == "warning"


def test_secret_is_encrypted_and_not_returned(database):
    with store.transaction() as s:
        response = save_provider(s, ProviderConfig(base_url="https://example.invalid/v1", model="test"), "private-key")
        assert "private-key" not in json.dumps(response)
        assert "private-key" not in s.query(AiSetting).get("runtime").secret
        assert provider(s)[1] == "private-key"


def test_provider_endpoint_change_requires_fresh_key_and_public_view_filters_unknown_fields(database, monkeypatch):
    from pydantic import ValidationError
    with store.transaction() as s:
        save_provider(s, ProviderConfig(base_url="https://first.example.invalid/v1", model="test"), "first-key")
    with pytest.raises(HTTPException) as conflict:
        with store.transaction() as s:
            save_provider(s, ProviderConfig(base_url="https://second.example.invalid/v1", model="test"))
    assert conflict.value.status_code == 409
    with store.transaction() as s:
        save_provider(s, ProviderConfig(base_url="https://second.example.invalid/v1", model="test"), "second-key")
        row = s.query(AiSetting).get("runtime")
        row.data = store.dumps({**json.loads(row.data), "api_key": "must-not-leak", "unexpected": "hidden"})
        view = api.public_settings(row)
        assert "api_key" not in view and "unexpected" not in view
        assert provider(s)[1] == "second-key"
    with pytest.raises(ValidationError):
        ProviderConfig(base_url="http://provider.example.invalid/v1", model="test")
    assert ProviderConfig(base_url="http://127.0.0.1:9000/v1", model="test").base_url.startswith("http://")
    monkeypatch.setattr("const.MOD", "release")
    with pytest.raises(ValidationError):
        ProviderConfig(base_url="http://127.0.0.1:9000/v1", model="test")


def test_provider_settings_update_uses_optimistic_version(database):
    app = FastAPI()
    app.include_router(api.router)
    app.dependency_overrides[api.cover_header] = lambda: ADMIN
    client = TestClient(app)
    body = {**model_config().dict(), "api_key": "test-key", "expected_version": 1}
    first = client.post("/auto-task/ai/settings/update", json=body)
    assert first.status_code == 200, first.text
    assert first.json()["data"]["version"] == 2
    stale = client.post("/auto-task/ai/settings/update", json={**body, "model": "stale"})
    assert stale.status_code == 409
    latest = client.post("/auto-task/ai/settings/update", json={**body, "model": "next", "expected_version": 2})
    assert latest.status_code == 200
    assert latest.json()["data"]["model"] == "next"


def test_model_http_request_never_follows_redirect_with_provider_authorization(database, monkeypatch):
    from auto_task.runtime import Context
    with store.transaction() as s:
        save_provider(s, model_config(), "test-only-key")
    enqueue()
    job = store.claim()
    options = {}

    class Response:
        ok = False
        status_code = 302

    monkeypatch.setattr("auto_task.runtime.requests.post",
                        lambda *args, **kwargs: options.update(kwargs) or Response())
    with pytest.raises(RuntimeError, match="HTTP 302"):
        Context(job).ask("test", {"title": "Example", "type": "object"})
    assert options["allow_redirects"] is False


def test_dedup_ignores_option_order_not_question_case():
    a = {"question_type": "single_choice", "original": "Value of x?", "options": [{"label": "A", "original": "one"}, {"label": "B", "original": "two"}]}
    b = {**a, "options": list(reversed(a["options"]))}
    assert canonical_question(a) == canonical_question(b)
    assert canonical_question(a) != canonical_question({**a, "original": "Value of X?"})


def test_extraction_worker_persists_questions_and_reuses_duplicate(database, monkeypatch):
    from auto_task.runtime import Context
    question = {"question_type": "short_answer", "language": "en", "original": "2+2?", "zh": "2+2?", "options": [], "answer": "4", "explanation": "2+2=4", "tags": ["arithmetic"]}
    monkeypatch.setattr(Context, "ask", lambda self, prompt, schema, images=(): {"decision": "normal", "reason": ""} if schema["title"] == "Moderation" else {"questions": [question, question]})
    job_id = enqueue()
    assert run_once()
    with store.transaction() as s:
        job = s.query(TaskJob).get(job_id)
        assert job.status == "success"
        assert s.query(AiResource).filter_by(kind="question").count() == 1
        ids = json.loads(job.result)["question_ids"]
        assert ids[0] == ids[1]


def test_question_dedup_never_revives_deleted_questions(database, monkeypatch):
    from types import SimpleNamespace
    from auto_task import handlers

    exact = {"question_type": "short_answer", "language": "en", "original": "2+2?", "zh": "2+2?",
             "options": [], "answer": "4", "explanation": "", "tags": []}
    similar = {**exact, "original": "What is 2 + 2?", "zh": "What is 2 + 2?"}
    with store.transaction() as s:
        collection = resource(s, "42", "collection", {"question_ids": []})
        scope = "owner:42"
        tombstone = AiResource(id="deleted-question", kind="question", owner_id="42", owner_username="student",
                               data=store.dumps(exact), search_text=exact["original"], deleted=1,
                               canonical_key=store.digest({"scope": scope, "question": canonical_question(exact)}))
        s.add(tombstone)
        s.flush()
        job = SimpleNamespace(resource_id=collection.id, owner_id="42", owner_username="student", id="exact-job",
                              task_type="question.extract", payload="{}")
        with pytest.raises(ValueError, match="explicit restoration required"):
            apply_result(s, job, {"questions": [dict(exact)]}, collection.version)

    ctx = SimpleNamespace(job={"course_id": None, "owner_id": "42"})
    monkeypatch.setattr(handlers, "ask", lambda *_args, **_kwargs: pytest.fail("deleted candidates must not reach approximate dedup"))
    handlers.match_duplicate(ctx, similar)
    assert "_duplicate" not in similar
    with store.transaction() as s:
        collection = s.query(AiResource).filter_by(kind="collection").one()
        job = SimpleNamespace(resource_id=collection.id, owner_id="42", owner_username="student", id="similar-job",
                              task_type="question.extract", payload="{}")
        result = apply_result(s, job, {"questions": [similar]}, collection.version)
        created_id = result["question_ids"][0]
        assert created_id != "deleted-question"

    with store.transaction() as s:
        assert s.query(AiResource).get("deleted-question").deleted == 1
        assert s.query(AiResource).get(created_id).deleted == 0


def test_resource_http_contract_and_idempotent_job(client):
    response = client.post("/auto-task/ai/resources/create", json={"kind": "collection", "data": {"text": "test"}})
    assert response.status_code == 200, response.text
    item = response.json()["data"]
    payload = {"task_type": "question.extract", "operation_key": "abc", "resource_id": item["id"]}
    a = client.post("/auto-task/ai/jobs/create", json=payload)
    b = client.post("/auto-task/ai/jobs/create", json=payload)
    assert a.status_code == 200, a.text
    assert a.json()["data"]["taskId"] == b.json()["data"]["taskId"]


def test_client_cannot_forge_candidate_or_identity(client):
    response = client.post("/auto-task/ai/resources/create", json={"kind": "question", "data": {"candidate": {"score": 100}}})
    assert response.status_code == 400
    with pytest.raises(HTTPException):
        store.principal({"username": "someone", "userId": "not-an-id"})


def test_resource_owner_isolation(client):
    with store.transaction() as s:
        resource(s, "99")
    assert client.get("/auto-task/ai/resources/res-99").status_code == 403


def test_suspension_only_leaves_account_status(client):
    with store.transaction() as s:
        row = store.account(s, "42")
        row.status, row.reason = "suspended", "review required"
    result = client.get("/auto-task/ai/account")
    assert result.status_code == 200
    assert "balance" not in result.json()["data"]
    assert client.get("/auto-task/ai/capabilities").status_code == 403


def test_media_idempotency_and_path_safety(client):
    a = client.post("/auto-task/ai/files/upload", files={"file": ("notes.md", b"# hello", "text/markdown")}, data={"sequence": "one"})
    b = client.post("/auto-task/ai/files/upload", files={"file": ("notes.md", b"# hello", "text/markdown")}, data={"sequence": "one"})
    assert a.status_code == 200, a.text
    assert a.json()["data"]["id"] == b.json()["data"]["id"]
    assert client.get("/auto-task/ai/files/" + a.json()["data"]["id"]).content == b"# hello"
    from auto_task.configuration import safe_path
    with pytest.raises(HTTPException):
        safe_path("../../etc/passwd")


def test_upload_media_type_is_derived_from_the_allowlisted_extension(client):
    response = client.post(
        "/auto-task/ai/files/upload",
        files={"file": ("notes.md", b"# trusted text", "image/png")},
    )
    assert response.status_code == 200, response.text
    with store.transaction() as s:
        saved = s.query(AiFile).get(response.json()["data"]["id"])
        assert saved.media_type == "text/markdown"


def test_exam_planning_capacity_and_no_accounts_created():
    result = plan_exam({"candidates": [{"user_id": str(x)} for x in range(7)], "rooms": [{"name": "A", "capacity": 5}, {"name": "B", "capacity": 3}, {"name": "C", "capacity": 1}]})
    assert len(result["rooms"]) == 2
    assert result["published"] is False
    assert sum(len(r["assigned"]) for r in result["rooms"]) == 7
    with pytest.raises(ValueError):
        plan_exam({"candidates": [{"user_id": "1"}], "rooms": []})


def test_mysql_concurrent_claims_respect_global_limit(database):
    if database.dialect.name != "mysql":
        pytest.skip("Real MySQL locking test")
    from concurrent.futures import ThreadPoolExecutor
    with store.transaction() as s:
        for number in range(10, 20):
            user = {"userId": number, "username": str(number), "roles": ["user"], "groups": []}
            row = resource(s, str(number))
            row.owner_username = user["username"]
            store.submit(s, user, "question.extract", {}, "concurrent", row)
    with ThreadPoolExecutor(max_workers=6) as pool:
        claims = list(pool.map(lambda _: store.claim(), range(6)))
    claimed = [row for row in claims if row]
    assert len(claimed) == 2
    assert len({row["id"] for row in claimed}) == 2


def test_long_question_input_survives_mysql_storage(database):
    enqueue(data={"text": "x" * 150000})
    with store.transaction() as s:
        assert len(json.loads(s.query(TaskJob).one().payload)["text"]) == 150000


def test_queued_job_waits_for_points_without_losing_payload(database):
    job_id = enqueue()
    with store.transaction() as s:
        store.ledger(s, "42", "drain", -1, "42", "usage")
    assert store.claim() is None
    with store.transaction() as s:
        assert s.query(TaskJob).get(job_id).status == "blocked"
        store.ledger(s, "42", "refill", 10, "7", "grant")
    assert store.claim(datetime.utcnow() + timedelta(minutes=1))["id"] == job_id


def test_student_cannot_delete_teacher_interview(client):
    with store.transaction() as s:
        row = resource(s, "7", "interview", {"student_id": "42", "status": "pending", "expires_at": (datetime.utcnow() + timedelta(hours=1)).isoformat()})
    assert client.get("/auto-task/ai/resources/res-7").status_code == 200
    assert client.post("/auto-task/ai/resources/res-7/delete").status_code == 403


def test_interview_identity_preflight_and_sequenced_answer(client, monkeypatch):
    from auto_task.runtime import Context
    monkeypatch.setattr(Context, "ask", lambda self, prompt, schema, images=():
        {"decision": "normal", "reason": ""} if schema["title"] == "Moderation" else
        {"question": "Explain your algorithm", "focus": "complexity", "complete": False, "summary": "", "score": 0, "confidence": 0.5})
    with store.transaction() as s:
        resource(s, "7", "interview", {"student_id": "42", "status": "pending", "turns": [], "expires_at": (datetime.utcnow() + timedelta(hours=1)).isoformat(), "max_questions": 3})
    preflight = {"operation_key": "start", "evidence": {"https": True, "camera": True, "microphone": True, "screen": True, "single_display": True}}
    response = client.post("/auto-task/ai/interviews/res-7/preflight", json=preflight)
    assert response.status_code == 200, response.text
    assert run_once()
    response = client.post("/auto-task/ai/interviews/res-7/answer", json={"operation_key": "answer-1", "sequence": 1, "text": "O(n)"})
    assert response.status_code == 200, response.text
    assert client.post("/auto-task/ai/interviews/res-7/answer", json={"operation_key": "wrong", "sequence": 0, "text": "x"}).status_code == 409


def test_document_index_and_grounded_rag(database, monkeypatch):
    from auto_task.runtime import Context
    monkeypatch.setattr(Context, "ask", lambda self, prompt, schema, images=():
        {"decision": "normal", "reason": ""} if schema["title"] == "Moderation" else
        {"answer": "A queue is FIFO", "citations": ["res-42:v1:text:1:0"]})
    with store.transaction() as s:
        doc = resource(s, "42", "document", {"text": "A queue is FIFO. A stack is LIFO."})
        store.submit(s, USER, "course.index", {"resource_version": 1}, "index", doc)
    assert run_once()
    with store.transaction() as s:
        assert s.query(AiResource).filter_by(kind="chunk").count() == 1
        conversation = resource(s, "43", "conversation", {})
        conversation.owner_id = "42"
        store.submit(s, USER, "course.ask", {"resource_version": 1, "message": "queue"}, "ask", conversation)
    assert run_once()
    with store.transaction() as s:
        job = s.query(TaskJob).filter_by(task_type="course.ask").one()
        assert job.status == "success"
        assert json.loads(job.result)["citations"] == ["res-42:v1:text:1:0"]
        detail = json.loads(job.result)["citation_details"][0]
        assert detail["citation"] == "res-42:v1:text:1:0"
        assert detail["resource_id"] == "res-42"
        assert detail["resource_version"] == detail["page"] == 1
        assert detail["file_id"] == "text" and detail["file_name"] is None and detail["offset"] == 0
        assert detail["source_sha256"] == hashlib.sha256(b"A queue is FIFO. A stack is LIFO.").hexdigest()
        assert detail["source"] == "course" and detail["evidence_confidence"] == "lexical_match"
        assert set(detail["lexical_overlap"]) >= {"queue", "fifo"}
        conversation = s.query(AiResource).filter_by(id="res-43").one()
        assistant = json.loads(conversation.data)["messages"][-1]
        assert assistant["citation_details"][0]["page"] == 1


def test_rag_joins_parent_and_excludes_stale_or_cross_owner_chunks(database):
    from auto_task.documents import course_answer
    from types import SimpleNamespace

    with store.transaction() as s:
        current = resource(s, "42", "document", {"title": "queues", "indexed_version": 2})
        current.version = 2
        stale = resource(s, "44", "document", {"title": "old", "indexed_version": 3})
        stale.version = 3
        other = resource(s, "99", "document", {"title": "private", "indexed_version": 2})
        other.version = 2
        for parent, text, source_version in ((current, "queue uses FIFO order", 1),
                                             (stale, "queue uses a stale LIFO claim", 1),
                                             (other, "queue is private", 1)):
            citation = "%s:v%d:text:1:0" % (parent.id, source_version)
            s.add(AiResource(kind="chunk", owner_id=parent.owner_id, owner_username="student",
                             parent_id=parent.id, data=store.dumps({"resource_id": parent.id,
                                 "resource_version": source_version, "source_sha256": "a" * 64,
                                 "page": 1, "offset": 0, "citation": citation, "text": text}),
                             search_text=text, state="ready"))

    captured = {}
    context = SimpleNamespace(job={"course_id": None, "owner_id": "42", "resource_id": "conversation"})
    context.ask = lambda prompt, *_: (captured.update(payload=json.loads(prompt.split("\n")[-1])) or
                                      {"answer": "The queue uses FIFO order.",
                                       "citations": ["res-42:v1:text:1:0"]})
    result = course_answer(context, {"messages": []}, {"message": "queue order"})
    assert result["citations"] == ["res-42:v1:text:1:0"]
    assert [item["citation"] for item in captured["payload"]["excerpts"]] == ["res-42:v1:text:1:0"]


def test_rag_flags_paraphrased_citation_for_review_and_rejects_duplicates(database):
    from auto_task.documents import course_answer
    from types import SimpleNamespace

    with store.transaction() as s:
        document = resource(s, "42", "document", {"indexed_version": 2})
        document.version = 2
        part = {"resource_id": document.id, "resource_version": 1, "source_sha256": "b" * 64,
                "page": 1, "offset": 0, "citation": "res-42:v1:text:1:0", "text": "Queue order is FIFO."}
        s.add(AiResource(kind="chunk", owner_id="42", owner_username="student", parent_id=document.id,
                         data=store.dumps(part), search_text=part["text"], state="ready"))
    context = SimpleNamespace(job={"course_id": None, "owner_id": "42", "resource_id": "conversation"})
    context.ask = lambda *_: {"answer": "Binary trees have logarithmic height.",
                              "citations": ["res-42:v1:text:1:0"]}
    result = course_answer(context, {}, {"message": "queue"})
    assert result["citation_details"][0]["evidence_confidence"] == "requires_review"
    assert result["citation_details"][0]["lexical_overlap"] == []
    context.ask = lambda *_: {"answer": "Queue order is FIFO.",
                              "citations": ["res-42:v1:text:1:0", "res-42:v1:text:1:0"]}
    with pytest.raises(ValueError, match="ungrounded"):
        course_answer(context, {}, {"message": "queue"})


def test_native_grade_adoption_checks_source_and_preserves_manual_score(client):
    from auto_task.legacy_bridge import answer_snapshot
    client.app.dependency_overrides[api.cover_header] = lambda: ADMIN
    with store.transaction() as s:
        ps = ProblemSet(psid=1, name="exam", type=1, groupInfo="[]", config="{}", username="teacher", groupId=1, tag="test")
        s.add(ps)
        s.add(ProblemSubjective(pid=1, type=1, config='{"judgeConfig":[{"name":"correct","score":10}]}', description="question", username="teacher"))
        s.add(ProblemSetAnswerSheet(asid=1, psid=1, username="student"))
        s.flush()
        s.add(ProblemSetAnswerSheetDetail(asd_id=1, asid=1, gid=1, pid=1, answer='["answer"]'))
        s.flush()
        _, snapshot = answer_snapshot(s, 1, 1, 1, "student")
        grading = resource(s, "7", "grading", {})
        grading.psid = 1
        job = store.submit(s, ADMIN, "subjective.grade", {}, "grade", grading)
        s.flush()
        job.status = "success"
        candidate = {"source": snapshot, "judgeLog": [{"name": "correct", "score": 10, "jScore": 8}], "judgeComment": "evidence"}
        grading.data = store.dumps({"candidate": candidate, "candidate_job_id": job.id})
        grading.state = "review"
    response = client.post("/auto-task/ai/grading/res-7/adopt", json={"version": 1, "reason": "checked"})
    assert response.status_code == 200, response.text
    with store.transaction() as s:
        detail = s.query(ProblemSetAnswerSheetDetail).get(1)
        assert json.loads(detail.judgeLog)[0]["jScore"] == 8
        assert detail.judgeLock_username == "teacher"
    replay = client.post("/auto-task/ai/grading/res-7/adopt", json={"version": 1, "reason": "retry"})
    assert replay.json()["data"]["replay"]


def test_plain_provider_call_records_cost(database, monkeypatch):
    from auto_task.runtime import Context
    with store.transaction() as s:
        save_provider(s, ProviderConfig(provider="chat_completions", base_url="https://example.invalid/v1", model="test",
                                       input_usd_per_million=1, output_usd_per_million=2), "private-key")
    job_id = enqueue()
    job = store.claim()
    monkeypatch.setattr(Context, "_http", lambda *args: ({"answer": "4"}, {"prompt_tokens": 1000, "completion_tokens": 500}))
    assert Context(job).ask("test", {}) == {"answer": "4"}
    with store.transaction() as s:
        assert store.account(s, "42").balance == Decimal("-2")
        assert s.query(AiLedger).filter_by(task_id=job_id).one().cost_usd == Decimal("0.002")


def test_sandbox_arguments_cannot_access_host_network(tmp_path):
    from auto_task.sandbox import arguments
    argv = arguments("trusted-image@sha256:abc", tmp_path, "task", ["python", "main.py"])
    assert argv[argv.index("--network") + 1] == "none"
    assert "--privileged" not in argv
    assert "ALL" == argv[argv.index("--cap-drop") + 1]
    assert "--read-only" in argv


def test_api_only_uses_existing_method_policy(client):
    assert all(set(route.methods or ()) <= {"GET", "HEAD", "POST"} for route in client.app.routes)


@pytest.mark.parametrize("task_type,kind,data", [
    ("question.solve", "question", {"original": "2+2", "zh": "2+2"}),
    ("question.verify", "question", {"original": "2+2", "answer": "4"}),
    ("question.chat", "question", {"original": "2+2"}),
    ("authoring.generate", "authoring", {"requirements": "Print 4"}),
    ("subjective.grade", "grading", {"question": "2+2", "answer": "4"}),
    ("plagiarism.scan", "plagiarism", {}),
    ("ide.execute", "ide", {"language": "python", "code": "print(4)"}),
    ("ide.tutor", "ide", {"language": "python", "code": "print(4)"}),
])
def test_workflow_handlers_apply_validated_results(database, monkeypatch, task_type, kind, data):
    monkeypatch.setattr("auto_task.access.fetch_actor", lambda subject, **_: USER)
    from auto_task.runtime import Context
    results = {
        "Moderation": {"decision": "normal", "reason": ""},
        "Answer": {"answer": "4", "explanation": "addition", "tags": ["arithmetic"]},
        "Verification": {"passed": True, "evidence": ["checked"], "corrected_answer": "4", "explanation": "checked"},
        "Feedback": {"answer": "explanation", "citations": []},
        "Authoring": {"title": "test", "description": "Print 4", "reference_code": "print(4)", "language": "python",
                      "tests": [{"input": "", "output": "4"}], "tags": [], "explanation": "constant"},
        "Grade": {"criteria": [{"name": "correct", "score": 10, "max_score": 10, "evidence": "4"}],
                  "confidence": 1, "summary": "correct", "needs_review": False},
    }
    monkeypatch.setattr(Context, "ask", lambda self, prompt, schema, images=(): results[schema["title"]])
    monkeypatch.setattr("auto_task.sandbox.execute", lambda *args: {"exit_code": 0, "output": "4", "language": "python"})
    monkeypatch.setattr("auto_task.plagiarism.scan", lambda *args: {"pairs": [], "needs_review": True})
    with store.transaction() as s:
        row = resource(s, "42", kind, data)
        job = store.submit(s, USER, task_type, {"resource_version": 1, "message": "why"}, "workflow", row)
        s.flush()
        job_id = job.id
    assert run_once()
    with store.transaction() as s:
        job = s.query(TaskJob).get(job_id)
        assert job.status == "success", job.error
        assert s.query(AiResource).get("res-42").version == 2
        assert job.result


def test_create_idempotency_survives_resource_advancement(client, monkeypatch):
    from auto_task.runtime import Context
    monkeypatch.setattr(Context, "ask", lambda self, prompt, schema, images=(): {"decision": "normal", "reason": ""} if schema["title"] == "Moderation" else {"questions": []})
    item = client.post("/auto-task/ai/resources/create", json={"kind": "collection", "data": {"text": "test"}}).json()["data"]
    body = {"task_type": "question.extract", "operation_key": "late-retry", "resource_id": item["id"]}
    first = client.post("/auto-task/ai/jobs/create", json=body).json()["data"]
    assert run_once()
    second = client.post("/auto-task/ai/jobs/create", json=body)
    assert second.status_code == 200
    assert second.json()["data"]["taskId"] == first["taskId"]


def test_native_publish_maps_option_labels_to_native_positions(client, monkeypatch):
    from server import auto_task_actions as actions
    monkeypatch.setattr(actions, "group_manager", lambda *args: None)
    with store.transaction() as s:
        s.add(ProblemGroup(gid=1, name="choices", type=0, problemInfo="[]", username="student"))
        resource(s, "42", "question", {"question_type": "single_choice", "original": "Q", "zh": "Q",
                                      "options": [{"label": "B", "zh": "yes"}, {"label": "A", "zh": "no"}]})
    body = {"version": 1, "group_id": 1, "score": 5, "correct_labels": ["B"], "dry_run": False}
    response = client.post("/auto-task/ai/questions/res-42/publish", json=body)
    assert response.status_code == 200, response.text
    with store.transaction() as s:
        assert json.loads(s.query(ProblemObjective).one().answer) == ["A"]
    assert client.post("/auto-task/ai/questions/res-42/publish", json=body).json()["data"]["replay"]
    assert client.post("/auto-task/ai/questions/res-42/publish", json={**body, "score": 6}).status_code == 409


def test_legacy_import_requires_operator_attribution(database):
    from auto_task.migrate import import_legacy
    with store.transaction() as s:
        s.add(AutoTaskRun(id="old", task_type="subjective_review", username="student", status="pending"))
        s.flush()
        s.add(AutoTaskRunLog(task_id="old", tag="payload", content='{"psid":1,"username":"student"}'))
    assert import_legacy({})["unresolved_task_ids"] == ["old"]
    with pytest.raises(ValueError):
        import_legacy({"old": ADMIN}, True)
    result = import_legacy({"old": ADMIN}, True, True)
    assert result["applied"]
    with store.transaction() as s:
        assert s.query(TaskJob).get("old").owner_id == "7"


def test_fulltext_and_literal_wildcard_search(client):
    with store.transaction() as s:
        resource(s, "42", "question", {"zh": "binary tree"})
        s.query(AiResource).get("res-42").search_text = "binary tree 100% coverage"
    response = client.post("/auto-task/ai/resources/list", json={"kind": "question", "keyword": "binary"})
    assert response.status_code == 200, response.text
    assert response.json()["data"]["totalNum"] == 1
    response = client.post("/auto-task/ai/resources/list", json={"kind": "question", "keyword": "zzzz%"})
    assert response.json()["data"]["totalNum"] == 0


def test_svg_rejects_script_external_reference_and_entities():
    from auto_task.assets import sanitize_svg
    assert b"svg" in sanitize_svg('<svg xmlns="http://www.w3.org/2000/svg"><path d="M0 0 L1 1"/></svg>')
    for value in ['<svg><script>alert(1)</script></svg>', '<svg><path onclick="alert(1)"/></svg>',
                  '<!DOCTYPE svg><svg/>', '<svg><image href="https://example.invalid"/></svg>']:
        with pytest.raises(ValueError):
            sanitize_svg(value)


def test_roster_header_blank_lines_and_duplicates():
    from auto_task.handlers import parse_roster
    assert parse_roster("\nstudentNo,name,username\n20260001,Alice,alice\n")[0]["username"] == "alice"
    with pytest.raises(ValueError):
        parse_roster("1,A,a\n1,B,b")


def test_opaque_operation_keys_are_case_sensitive(database):
    assert enqueue("Case") != enqueue("case")


def test_legacy_endpoints_respect_ai_suspension(client):
    from server.auto_task import router as legacy
    client.app.include_router(legacy)
    with store.transaction() as s:
        store.account(s, "42").status = "suspended"
    assert client.post("/auto-task/list", json={}).status_code == 403


def test_sse_finishes_after_terminal_events(client):
    task_id = enqueue()
    job = store.claim()
    store.complete(task_id, job["lease_token"], {"ok": True})
    response = client.get("/auto-task/ai/jobs/" + task_id + "/events")
    assert response.status_code == 200
    assert "text/event-stream" in response.headers["content-type"]
    assert '"tag":"result"' in response.text


def test_job_views_include_lifecycle_and_review_state(client):
    with store.transaction() as s:
        item = resource(s, "42", "grading", {"title": "candidate"})
        job = store.submit(s, USER, "subjective.grade", {"resource_version": item.version}, "view", item)
        job.status, job.attempt, item.state = "success", 1, "review"
        job.finished_at = datetime.utcnow()
        task_id = job.id
    detail = client.get("/auto-task/ai/jobs/" + task_id).json()["data"]
    assert detail["review_required"] is True
    assert detail["max_attempts"] == 3
    assert detail["created_at"].endswith("Z") and detail["finished_at"].endswith("Z")
    listed = client.post("/auto-task/ai/jobs/list", json={"pageNow": 1, "pageSize": 20}).json()["data"]["rows"]
    assert listed[0]["review_required"] is True


def test_job_list_hides_tasks_of_deleted_resources(client):
    """B-01: list must not surface jobs whose detail 404s after resource deletion."""
    with store.transaction() as s:
        item = resource(s, "42", "grading", {"title": "deleted-parent"})
        job = store.submit(s, USER, "subjective.grade", {"resource_version": item.version}, "gone", item)
        job.status = "success"
        job.finished_at = datetime.utcnow()
        task_id = job.id
        resource_id = item.id
        s.query(AiResource).filter_by(id=resource_id).update({"deleted": 1})

    listed = client.post("/auto-task/ai/jobs/list", json={"pageNow": 1, "pageSize": 50}).json()["data"]["rows"]
    assert all(row["id"] != task_id for row in listed)
    detail = client.get("/auto-task/ai/jobs/" + task_id)
    assert detail.status_code == 404


def test_queue_scans_past_blocked_first_hundred(database):
    with store.transaction() as s:
        first = resource(s, "42")
        store.submit(s, USER, "question.extract", {}, "active", first)
    assert store.claim()
    with store.transaction() as s:
        for index in range(110):
            store.submit(s, USER, "question.extract", {}, "blocked-" + str(index), first)
        other = {"userId": 99, "username": "other"}
        row = resource(s, "99")
        row.owner_username = "other"
        expected = store.submit(s, other, "question.extract", {}, "eligible", row).id
    assert store.claim()["id"] == expected


@pytest.mark.parametrize("kind", ["objective", "subjective"])
def test_native_import_preserves_oj_standard(client, monkeypatch, kind):
    from server import auto_task_native as native
    monkeypatch.setattr(native, "group_manager", lambda *args: None)
    with store.transaction() as s:
        s.add(ProblemGroup(gid=1, name="native", type=0 if kind == "objective" else 1,
                           problemInfo='[{"pid":1}]', username="student"))
        if kind == "objective":
            s.add(ProblemObjective(pid=1, type=0, content='{"description":"2+2?","choice":["4","5"]}',
                                  answer='["A"]', username="student"))
        else:
            s.add(ProblemSubjective(pid=1, type=1, config='{"judgeConfig":[{"name":"correct","score":10,"answer":"4"}]}',
                                   description="2+2?", username="student"))
    body = {"group_id": 1, "pid": 1}
    result = client.post("/auto-task/ai/native/import", json=body)
    assert result.status_code == 200, result.text
    value = result.json()["data"]
    assert value["data"]["native"]["type"] == kind
    assert value["data"]["answer"] == ("A" if kind == "objective" else "4")
    again = client.post("/auto-task/ai/native/import", json=body)
    assert again.json()["data"]["id"] == value["id"]
    assert again.json()["data"]["replay"]
    assert client.post("/auto-task/ai/native/import", json={"group_id": 1, "pid": 999}).status_code == 403


def test_native_import_rejects_deleted_replay_without_restoring(client, monkeypatch):
    from server import auto_task_native as native
    monkeypatch.setattr(native, "group_manager", lambda *args: None)
    with store.transaction() as s:
        s.add(ProblemGroup(gid=1, name="native", type=0, problemInfo='[{"pid":1}]', username="student"))
        s.add(ProblemObjective(pid=1, type=0, content='{"description":"2+2?","choice":["4","5"]}',
                              answer='["A"]', username="student"))
    body = {"group_id": 1, "pid": 1}
    imported = client.post("/auto-task/ai/native/import", json=body).json()["data"]
    assert client.post("/auto-task/ai/resources/%s/delete" % imported["id"]).status_code == 200
    response = client.post("/auto-task/ai/native/import", json=body)
    assert response.status_code == 409
    with store.transaction() as s:
        assert s.query(AiResource).get(imported["id"]).deleted == 1


def test_native_import_refreshes_changed_source_with_explicit_version(client, monkeypatch):
    from server import auto_task_native as native
    monkeypatch.setattr(native, "group_manager", lambda *args: None)
    with store.transaction() as s:
        s.add(ProblemGroup(gid=1, name="native", type=0,
                           problemInfo='[{"pid":1}]', username="student"))
        s.add(ProblemObjective(pid=1, type=0,
                               content='{"description":"old","choice":["A","B"]}',
                               answer='["A"]', username="student"))
    imported = client.post("/auto-task/ai/native/import", json={"group_id": 1, "pid": 1}).json()["data"]
    annotated = client.post("/auto-task/ai/resources/%s/update" % imported["id"], json={
        "version": imported["version"], "data": {"tags": ["teacher-note"], "explanation": "keep this"},
    })
    assert annotated.status_code == 200, annotated.text
    imported = annotated.json()["data"]
    with store.transaction() as s:
        source = s.query(ProblemObjective).get(1)
        source.content = '{"description":"new","choice":["A","B"]}'
        source.answer = '["B"]'
    stale = client.post("/auto-task/ai/native/import", json={
        "group_id": 1, "pid": 1, "refresh_resource_id": imported["id"],
        "version": imported["version"] + 1,
    })
    assert stale.status_code == 409
    refreshed = client.post("/auto-task/ai/native/import", json={
        "group_id": 1, "pid": 1, "refresh_resource_id": imported["id"],
        "version": imported["version"],
    })
    assert refreshed.status_code == 200, refreshed.text
    value = refreshed.json()["data"]
    assert value["refreshed"] is True and value["version"] == imported["version"] + 1
    assert value["data"]["zh"] == "new" and value["data"]["answer"] == "B"
    assert value["data"]["tags"] == ["teacher-note"] and value["data"]["explanation"] == "keep this"
    replay = client.post("/auto-task/ai/native/import", json={
        "group_id": 1, "pid": 1, "refresh_resource_id": imported["id"],
        "version": value["version"],
    }).json()["data"]
    assert replay["replay"] is True and replay["refreshed"] is False


def test_native_import_identity_survives_local_annotation(client, monkeypatch):
    from server import auto_task_native as native
    monkeypatch.setattr(native, "group_manager", lambda *args: None)
    with store.transaction() as s:
        s.add(ProblemGroup(gid=1, name="native", type=0,
                           problemInfo='[{"pid":1}]', username="student"))
        s.add(ProblemObjective(pid=1, type=0,
                               content='{"description":"Q","choice":["A","B"]}',
                               answer='["A"]', username="student"))
    imported = client.post("/auto-task/ai/native/import", json={"group_id": 1, "pid": 1}).json()["data"]
    annotated = client.post("/auto-task/ai/resources/%s/update" % imported["id"], json={
        "version": imported["version"], "data": {"tags": ["local-tag"]},
    })
    assert annotated.status_code == 200, annotated.text
    repeated = client.post("/auto-task/ai/native/import", json={"group_id": 1, "pid": 1})
    assert repeated.status_code == 200, repeated.text
    assert repeated.json()["data"]["id"] == imported["id"]
    with store.transaction() as s:
        assert s.query(AiResource).filter_by(kind="question").count() == 1


def test_program_plan_is_native_private_and_does_not_claim_judging(client):
    client.app.dependency_overrides[api.cover_header] = lambda: ADMIN
    with store.transaction() as s:
        resource(s, "7", "question", {"question_type": "programming", "zh": "# A+B"})
    response = client.post("/auto-task/ai/native/program/res-7/plan", json={"version": 1, "title": "AI test"})
    assert response.status_code == 200, response.text
    plan = response.json()["data"]
    assert plan["problem"]["isPublic"] == 0
    assert plan["problem"]["source"] == "autolab-ai:res-7:v1"
    assert plan["description"]["markdownDescription"] == "# A+B"
    assert not plan["ready_for_judging"]


def test_pptx_slides_are_ordered_and_not_flattened(tmp_path):
    import zipfile
    from types import SimpleNamespace
    from auto_task.documents import extract_document
    path = tmp_path / "slides.pptx"
    with zipfile.ZipFile(path, "w") as archive:
        for number in [10, 2, 1]:
            archive.writestr("ppt/slides/slide%d.xml" % number,
                            '<p xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:t>Slide %d</a:t></p>' % number)
    pages = extract_document(path, SimpleNamespace(check=lambda: None))
    assert [page["page"] for page in pages] == [1, 2, 3]
    assert pages[-1]["text"] == "Slide 10"


def test_pptx_uses_presentation_relationship_order(tmp_path):
    import zipfile
    from types import SimpleNamespace
    from auto_task.documents import extract_document

    path = tmp_path / "ordered.pptx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("ppt/slides/slide10.xml", '<p xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:t>Deck first</a:t></p>')
        archive.writestr("ppt/slides/slide2.xml", '<p xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:t>Deck second</a:t></p>')
        archive.writestr("ppt/presentation.xml", '<p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><p:sldIdLst><p:sldId id="256" r:id="rId1"/><p:sldId id="257" r:id="rId2"/></p:sldIdLst></p:presentation>')
        archive.writestr("ppt/_rels/presentation.xml.rels", '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Target="slides/slide10.xml"/><Relationship Id="rId2" Target="slides/slide2.xml"/></Relationships>')
    pages = extract_document(path, SimpleNamespace(check=lambda: None))
    assert [page["text"] for page in pages] == ["Deck first", "Deck second"]
    assert [page["page"] for page in pages] == [1, 2]


def test_docx_explicit_page_breaks_preserve_logical_page_numbers(tmp_path):
    from docx import Document
    from types import SimpleNamespace
    from auto_task.documents import extract_document

    path = tmp_path / "lecture.docx"
    document = Document()
    document.add_paragraph("Page one")
    document.add_page_break()
    document.add_paragraph("Page two")
    document.add_table(rows=1, cols=2).rows[0].cells[0].text = "A"
    document.tables[0].rows[0].cells[1].text = "B"
    document.save(path)
    pages = extract_document(path, SimpleNamespace(check=lambda: None))
    assert [page["page"] for page in pages] == [1, 2]
    assert "Page one" in pages[0]["text"]
    assert "Page two" in pages[1]["text"]
    assert "A | B" in pages[1]["text"]


def test_indexed_file_citation_has_name_hash_version_and_page(database, tmp_path):
    from types import SimpleNamespace
    from auto_task.documents import index_document

    path = tmp_path / "slides.txt"
    path.write_text("Dijkstra finds shortest paths", encoding="utf8")
    digest_value = hashlib.sha256(path.read_bytes()).hexdigest()
    with store.transaction() as s:
        s.add(AiFile(id="file-1", owner_id="42", filename="Lecture 1.txt", media_type="text/plain",
                     path="uploads/file-1.txt", sha256=digest_value, size=path.stat().st_size))
    context = SimpleNamespace(job={"resource_id": "doc-1", "payload": '{"resource_version":4}'},
                              check=lambda: None)
    chunk = index_document(context, [("file-1", path, "text/plain")], {})["chunks"][0]
    assert chunk["citation"] == "doc-1:v4:file-1:1:0"
    assert chunk["resource_id"] == "doc-1"
    assert chunk["resource_version"] == 4
    assert chunk["file_name"] == "Lecture 1.txt"
    assert chunk["source_sha256"] == digest_value


def test_document_chunks_endpoint_returns_only_current_version_with_locations(client):
    with store.transaction() as s:
        document = resource(s, "42", "document", {"title": "notes", "indexed_version": 2})
        document.version = 2
        current = {"citation": "res-42:v1:text:1:0", "resource_id": document.id,
                   "resource_version": 1, "file_id": "text", "file_name": None,
                   "page": 1, "offset": 0, "source_sha256": "c" * 64, "text": "Current"}
        stale = {**current, "citation": "res-42:v0:text:1:0", "resource_version": 0, "text": "Stale"}
        for part in (current, stale):
            s.add(AiResource(kind="chunk", owner_id="42", owner_username="student", parent_id=document.id,
                             data=store.dumps(part), search_text=part["text"], state="ready"))
    response = client.get("/auto-task/ai/resources/res-42/chunks?pageNow=1&pageSize=1")
    assert response.status_code == 200, response.text
    payload = response.json()["data"]
    assert payload["totalNum"] == 1
    assert payload["rows"][0]["citation"] == "res-42:v1:text:1:0"
    assert payload["rows"][0]["source_sha256"] == "c" * 64

    with store.transaction() as s:
        s.query(AiResource).filter_by(id="res-42").one().version = 3
    response = client.get("/auto-task/ai/resources/res-42/chunks")
    assert response.status_code == 200
    assert response.json()["data"]["totalNum"] == 0


def test_citation_resolver_checks_current_parent_version(client):
    with store.transaction() as s:
        document = resource(s, "42", "document", {"indexed_version": 2})
        document.version = 2
        part = {"citation": "res-42:v1:file-1:3:1200", "resource_id": document.id,
                "resource_version": 1, "file_id": "file-1", "file_name": "lecture.pdf",
                "page": 3, "offset": 1200, "source_sha256": "d" * 64, "text": "Source excerpt"}
        s.add(AiResource(kind="chunk", owner_id="42", owner_username="student", parent_id=document.id,
                         data=store.dumps(part), search_text=part["text"], state="ready"))
    response = client.get("/auto-task/ai/citations/resolve", params={"citation": part["citation"]})
    assert response.status_code == 200, response.text
    assert response.json()["data"]["file_name"] == "lecture.pdf"
    with store.transaction() as s:
        s.query(AiResource).filter_by(id="res-42").one().version = 3
    response = client.get("/auto-task/ai/citations/resolve", params={"citation": part["citation"]})
    assert response.status_code == 409


def test_empty_document_cannot_be_indexed():
    from types import SimpleNamespace
    from auto_task.documents import index_document
    with pytest.raises(ValueError):
        index_document(SimpleNamespace(job={"resource_id": "r", "payload": "{}"}), [], {})


def test_task_input_is_immutable_and_verified(database):
    from auto_task.entities import TaskInput
    job_id = enqueue()
    with store.transaction() as session:
        saved = session.query(TaskInput).filter_by(task_id=job_id).one()
        assert json.loads(saved.content)["text"] == "2+2?"
        session.query(AiResource).get("res-42").data = '{"text":"modified"}'
    claimed = store.claim()
    assert claimed["input_snapshot"]["data"]["text"] == "2+2?"


def test_interview_completion_uses_rubric_not_model_total(database, monkeypatch):
    from auto_task.runtime import Context
    from auto_task.handlers import handle
    payload = {"student_id": "42", "status": "running",
               "expires_at": (datetime.utcnow() + timedelta(hours=1)).isoformat(),
               "rubric": [{"name": "understanding", "max_score": 20}],
               "turns": [{"question": "why?", "answer": "because"}], "max_questions": 1}
    with store.transaction() as session:
        row = resource(session, "42", "interview", payload)
        store.submit(session, USER, "interview.advance", {"resource_version": 1}, "finish", row)
    def answer(self, prompt, schema, images=()):
        if schema["title"] == "Moderation":
            return {"decision": "normal", "reason": ""}
        return {"question": "", "focus": "", "complete": True, "summary": "reviewed",
                "score": 100, "confidence": 0.6, "criteria": [{"name": "understanding", "score": 10, "max_score": 20, "evidence": "answer"}]}
    monkeypatch.setattr(Context, "ask", answer)
    assert run_once()
    with store.transaction() as session:
        data = json.loads(session.query(AiResource).get("res-42").data)
        assert len(data["turns"]) == 1
        assert data["candidate"]["score"] == 50
        assert data["candidate"]["integrity_conclusion"] == "review"
        assert data["status"] == "completed"


def test_native_source_removal_revokes_import_access(client, monkeypatch):
    from server import auto_task_native as native
    monkeypatch.setattr(native, "group_manager", lambda *args: None)
    with store.transaction() as session:
        session.add(ProblemGroup(gid=1, name="native", type=0, problemInfo='[{"pid":1}]', username="student"))
        session.add(ProblemObjective(pid=1, type=0, content='{"description":"Q","choice":["A","B"]}', answer='["A"]', username="student"))
    result = client.post("/auto-task/ai/native/import", json={"group_id": 1, "pid": 1}).json()["data"]
    with store.transaction() as session:
        session.query(ProblemGroup).get(1).problemInfo = "[]"
    assert client.get("/auto-task/ai/resources/" + result["id"]).status_code == 403


def test_corrupt_snapshot_does_not_poison_next_job(database):
    from auto_task.entities import TaskInput
    first = enqueue("corrupt")
    second = enqueue("healthy")
    with store.transaction() as s:
        s.query(TaskInput).get(first).content = "invalid-json"
    assert store.claim()["id"] == second
    with store.transaction() as s:
        assert s.query(TaskJob).get(first).status == "failed"


def model_config():
    return ProviderConfig(provider="chat_completions", base_url="https://example.invalid/v1", model="test",
                          input_usd_per_million=1, cached_usd_per_million=0.5, output_usd_per_million=2)


def test_crashed_model_intent_blocks_until_idempotent_reconciliation(database):
    from auto_task.accounting import begin_run, settle_run
    from auto_task.entities import AiRun
    job_id = enqueue()
    job = store.claim()
    run_id = begin_run(job, model_config())
    with store.transaction() as s:
        assert s.query(AiRun).get(run_id).status == "running"
        s.query(TaskJob).get(job_id).lease_until = datetime.utcnow() - timedelta(seconds=1)
    assert store.claim() is None
    with store.transaction() as s:
        assert s.query(TaskJob).get(job_id).status == "blocked"
        assert s.query(AiRun).get(run_id).status == "unsettled"
    with pytest.raises(HTTPException) as error:
        enqueue("bypass")
    assert error.value.status_code == 409
    usage = {"input_tokens": 1000, "cached_input_tokens": 200, "output_tokens": 500}
    review = {**usage, "reason": "Provider statement verified", "actor_id": "7"}
    result = settle_run(run_id, usage, "unknown", reconciliation=review)
    assert Decimal(result["cost_usd"]) == Decimal("0.0019")
    assert settle_run(run_id, usage, "unknown", reconciliation=review)["replay"]
    with pytest.raises(HTTPException) as conflict:
        settle_run(run_id, usage, "unknown", reconciliation={**review, "reason": "different"})
    assert conflict.value.status_code == 409
    settle_run(run_id, {"input_tokens": 100000, "output_tokens": 500}, "success")
    with store.transaction() as s:
        assert s.query(AiLedger).filter_by(task_id=job_id).count() == 1
        assert s.query(AiRun).get(run_id).status == "reconciled"
        store.ledger(s, "42", "top-up", 100, "7", "dev credit")
    assert store.claim(datetime.utcnow() + timedelta(seconds=31))["id"] == job_id


@pytest.mark.parametrize("usage", [{}, {"input_tokens": 10}, {"input_tokens": 10, "output_tokens": -1},
                                  {"input_tokens": 10, "cached_input_tokens": 11, "output_tokens": 1}])
def test_unknown_usage_is_never_treated_as_free(database, usage):
    from auto_task.accounting import begin_run, settle_run
    enqueue()
    run_id = begin_run(store.claim(), model_config())
    assert settle_run(run_id, usage, "failed")["status"] == "unsettled"
    with store.transaction() as s:
        assert s.query(AiLedger).count() == 0
        assert store.has_unsettled_usage(s, "42")


def test_actual_call_sees_persisted_intent_and_settles_failure(database, monkeypatch):
    from auto_task.runtime import Context
    from auto_task.entities import AiRun
    with store.transaction() as s:
        save_provider(s, model_config(), "test-only-key")
    enqueue()
    job = store.claim()
    def call(ctx, *_):
        with store.transaction() as s:
            assert s.query(AiRun).filter_by(task_id=job["id"], status="running").count() == 1
        ctx.observed_usage = {"input_tokens": 10, "output_tokens": 20}
        raise ValueError("Invalid model JSON")
    monkeypatch.setattr(Context, "_http", call)
    with pytest.raises(ValueError, match="Invalid model JSON"):
        Context(job).ask("test", {})
    with store.transaction() as s:
        assert s.query(AiRun).one().status == "failed"
        assert s.query(AiLedger).one().cost_usd == Decimal("0.00005")


def test_reconcile_endpoint_requires_admin(client):
    response = client.get("/auto-task/ai/admin/usage/unsettled")
    assert response.status_code == 403
    assert client.post("/auto-task/ai/admin/usage/unknown/reconcile", json={"input_tokens": 0, "output_tokens": 0, "reason": "test"}).status_code == 403


def test_admin_usage_and_review_views_include_operational_context(client):
    from auto_task.accounting import begin_run
    task_id = enqueue()
    run_id = begin_run(store.claim(), model_config())
    with store.transaction() as s:
        s.query(AiRun).filter_by(id=run_id).update({"status": "unsettled"})
        s.add(AiReview(subject_id="42", task_id=task_id, decision="warning", reason="review"))
        store.account(s, "42").status = "warning"
    client.app.dependency_overrides[api.cover_header] = lambda: ADMIN
    usage = client.get("/auto-task/ai/admin/usage/unsettled").json()["data"]["rows"][0]
    assert usage["payer_id"] == "42" and usage["created_at"].endswith("Z")
    review = client.get("/auto-task/ai/reviews").json()["data"][0]
    assert review["account_status"] == "warning" and review["created_at"].endswith("Z")


def test_revision_history_survives_edits_delete_and_rollback(client):
    from auto_task.entities import AiRevision
    item = client.post("/auto-task/ai/resources/create", json={"kind": "collection", "data": {"text": "original"}}).json()["data"]
    path = "/auto-task/ai/resources/" + item["id"]
    assert client.post(path + "/update", json={"version": 1, "data": {"text": "edited"}}).status_code == 200
    with pytest.raises(RuntimeError):
        with store.transaction() as s:
            s.query(AiResource).get(item["id"]).data = '{"text":"rolled back"}'
            s.flush()
            raise RuntimeError()
    assert client.post(path + "/delete").status_code == 200
    history = client.get(path + "/history").json()["data"]
    assert history["totalNum"] == 3
    assert {entry["data"]["text"] for entry in history["rows"]} == {"original", "edited"}
    assert history["rows"][0]["deleted"]
    assert all(entry["actor_id"] == "42" for entry in history["rows"])
    with store.transaction() as s:
        assert s.query(AiRevision).count() == 3


def test_revision_history_requires_resource_access(client):
    with store.transaction() as s:
        row = resource(s, "99")
    assert client.get("/auto-task/ai/resources/" + row.id + "/history").status_code == 403


def test_third_party_content_does_not_suspend_teacher(database, monkeypatch):
    from auto_task.handlers import moderation
    from auto_task.runtime import Context
    enqueue()
    job = store.claim()
    monkeypatch.setattr(Context, "ask", lambda *args: {"decision": "violation", "reason": "student prompt injection"})
    moderation(Context(job), {"answer": "untrusted"}, [], third_party=True)
    with store.transaction() as s:
        assert store.account(s, "42").status == "active"
        assert s.query(AiReview).count() == 0
        event = s.query(AutoTaskRunLog).filter_by(task_id=job["id"], tag="log").order_by(AutoTaskRunLog.log_id.desc()).first()
        assert json.loads(event.content)["account_action"] == "none"


def test_content_moderation_replay_is_idempotent(database, monkeypatch):
    from auto_task.handlers import moderation
    from auto_task.runtime import Context
    from auto_task.entities import AiReview
    enqueue()
    job = store.claim()
    answers = iter(({"decision": "warning", "reason": "needs review"},
                    {"decision": "violation", "reason": "must not replace"}))
    monkeypatch.setattr(Context, "ask", lambda *args: next(answers))
    moderation(Context(job), {"text": "one"}, [])
    # A retry of the same claimed task must not call the model or add another
    # review. The lease remains valid in this fixture, as it would after a
    # worker-level retry/recovery path.
    moderation(Context(job), {"text": "one"}, [])
    with store.transaction() as s:
        rows = s.query(AiReview).filter_by(task_id=job["id"]).all()
        assert len(rows) == 1 and rows[0].decision == "warning"


def test_visible_ide_cases_capture_mismatch_and_runtime_error(database, monkeypatch):
    monkeypatch.setattr("auto_task.access.fetch_actor", lambda subject, **_: USER)
    with store.transaction() as s:
        row = resource(s, "42", "ide", {"language": "python", "code": "print(4)"})
        store.submit(s, USER, "ide.execute", {"tests": [{"input": "1", "output": "4"}, {"input": "2", "output": "5"}, {"input": "3", "output": "4"}]}, "cases", row)
    monkeypatch.setattr("auto_task.sandbox.execute", lambda ctx, lang, code, stdin: {"output": "4\n", "exit_code": 1 if stdin == "3" else 0})
    assert run_once()
    with store.transaction() as s:
        result = json.loads(s.query(TaskJob).one().result)
        assert result["scope"] == "visible_tests" and result["passed"] == 1 and result["total"] == 3


def test_model_purpose_preserves_its_own_price_snapshot(database):
    config = model_config()
    config = ProviderConfig.parse_obj({**config.dict(), "model_overrides": {"moderation": {
        "model": "review-model", "input_usd_per_million": 3, "cached_usd_per_million": 1, "output_usd_per_million": 4}}})
    with store.transaction() as s:
        save_provider(s, config, "test-only-key")
    with store.transaction() as s:
        selected, _ = provider(s, "moderation")
        assert selected.model == "review-model" and selected.input_usd_per_million == 3
        assert provider(s, "authoring")[0].model == "test"
    from pydantic import ValidationError
    with pytest.raises(ValidationError):
        ProviderConfig.parse_obj({**config.dict(), "model_overrides": {"moderation": {"model": "unpriced"}}})


def test_cumulative_file_quota_removes_rejected_upload(client):
    from auto_task.entities import AiFile
    from auto_task.configuration import data_root
    with store.transaction() as s:
        s.query(AiSetting).get("runtime").data = '{"max_storage_mb":1}'
        s.add(AiFile(owner_id="42", filename="old.txt", media_type="text/plain", path="old.txt", sha256="0"*64, size=1024*1024))
    response = client.post("/auto-task/ai/files/upload", files={"file": ("new.txt", b"test", "text/plain")})
    assert response.status_code == 413
    assert list((data_root() / "uploads").iterdir()) == []
    result = client.get("/auto-task/ai/files/storage").json()["data"]
    assert result["used_bytes"] == result["quota_bytes"]


def test_native_indeterminate_choice_roundtrip_keeps_type(client, monkeypatch):
    from server import auto_task_native as native, auto_task_actions as actions
    monkeypatch.setattr(native, "group_manager", lambda *args: None)
    monkeypatch.setattr(actions, "group_manager", lambda *args: None)
    with store.transaction() as s:
        s.add(ProblemGroup(gid=1, name="native", type=0, problemInfo='[{"pid":1}]', username="student"))
        s.add(ProblemObjective(pid=1, type=2, content='{"description":"Q","choice":["A","B"]}', answer='["A"]', username="student"))
    imported = client.post("/auto-task/ai/native/import", json={"group_id": 1, "pid": 1}).json()["data"]
    published = client.post("/auto-task/ai/questions/" + imported["id"] + "/publish", json={"version": 1, "group_id": 1, "correct_labels": ["A"], "score": 10, "dry_run": False})
    assert published.status_code == 200, published.text
    with store.transaction() as s:
        assert s.query(ProblemObjective).get(published.json()["data"]["pid"]).type == 2


def test_extraction_adds_to_collection_without_erasing_prior_questions(database):
    from types import SimpleNamespace
    with store.transaction() as s:
        row = resource(s, "42", "collection", {"question_ids": ["previous-question"]})
        job = SimpleNamespace(resource_id=row.id, owner_id="42", owner_username="student", id="test-job", task_type="question.extract")
        result = apply_result(s, job, {"questions": []}, 1)
        assert result["needs_input"] and row.state == "needs_input"
        assert json.loads(row.data)["question_ids"] == ["previous-question"]


def test_collection_manifest_batches_are_numbered_and_reversible(database):
    from auto_task.question_workflows import append_batch, detach_duplicate, detach_item, ensure_manifest, restore_duplicate, restore_item
    data = {"question_ids": [], "items": [], "batches": [], "next_question_number": 1}
    records = [{"question_id": "q1", "original": {"question_type": "short_answer", "zh": "one"}},
               {"question_id": "q2", "duplicate_of": "q-old", "original": {"question_type": "short_answer", "zh": "two"}}]
    appended = append_batch(data, "batch-a", records, {"text_sha256": "x"})
    assert [item["number"] for item in appended] == [1, 2]
    assert data["question_ids"] == ["q1", "q2"]
    assert data["next_question_number"] == 3
    with pytest.raises(ValueError, match="already applied"):
        append_batch(data, "batch-a", records)
    detach_item(data, 1)
    assert data["question_ids"] == ["q2"]
    restore_item(data, 1)
    assert data["question_ids"] == ["q1", "q2"]
    detach_duplicate(data, 2, "q-split")
    assert data["question_ids"] == ["q1", "q-split"]
    restore_duplicate(data, 2)
    assert data["question_ids"] == ["q1", "q-old"]


def test_collection_relinking_never_revives_deleted_materialization(client):
    original = {"question_type": "short_answer", "language": "zh", "original": "one", "zh": "一",
                "options": [], "answer": "一", "explanation": "", "tags": []}
    data = {"question_ids": ["canonical-question"], "items": [{
        "number": 1, "question_id": "canonical-question", "status": "active", "batch_id": "batch",
        "original": original, "duplicate_of": "canonical-question",
    }], "batches": [{"id": "batch", "source": {"file_ids": []}}], "next_question_number": 2}
    with store.transaction() as s:
        collection = resource(s, "42", "collection", data)
        s.add(AiResource(id="canonical-question", kind="question", owner_id="42", owner_username="student",
                         data=store.dumps(original), search_text="one", state="ready"))

    item_path = "/auto-task/ai/question-workflows/collections/%s/items/1" % collection.id
    split = client.post(item_path + "/detach-duplicate")
    assert split.status_code == 200, split.text
    materialized_id = split.json()["data"]["question"]["id"]
    assert client.post("/auto-task/ai/resources/%s/delete" % materialized_id).status_code == 200

    assert client.post(item_path + "/detach").status_code == 200
    assert client.post(item_path + "/restore").status_code == 409
    assert client.post(item_path + "/restore-duplicate").status_code == 200
    assert client.post(item_path + "/detach-duplicate").status_code == 409

    with store.transaction() as s:
        assert s.query(AiResource).get(materialized_id).deleted == 1
        manifest = json.loads(s.query(AiResource).get(collection.id).data)
        assert manifest["question_ids"] == ["canonical-question"]


def test_question_intake_enqueues_configurable_safe_pipeline(client):
    response = client.post("/auto-task/ai/question-workflows/intake", json={
        "title": "Batch one", "text": "1. What is FIFO?", "batch_id": "batch-1",
        "auto_solve": True, "auto_verify": True,
    })
    assert response.status_code == 200, response.text
    payload = response.json()["data"]
    assert payload["pipeline"] == {"auto_solve": True, "auto_verify": True}
    assert payload["batch_id"] == "batch-1"
    with store.transaction() as s:
        job = s.query(TaskJob).get(payload["taskId"])
        submitted = json.loads(job.payload)
        assert submitted["allow_execution"] is False
        assert submitted["pipeline"]["auto_verify"] is True
        assert s.query(AiResource).get(payload["collection"]["id"]).kind == "collection"


def test_question_intake_rejects_unsupported_media_before_enqueuing(client):
    uploaded = client.post(
        "/auto-task/ai/files/upload",
        files={"file": ("recording.mp4", b"not parsed by question extraction", "video/mp4")},
    )
    assert uploaded.status_code == 200, uploaded.text
    response = client.post("/auto-task/ai/question-workflows/intake", json={
        "title": "Unsupported source", "file_ids": [uploaded.json()["data"]["id"]],
    })
    assert response.status_code == 400
    with store.transaction() as s:
        assert s.query(TaskJob).count() == 0


def test_question_intake_rejects_a_full_collection_before_model_cost(client):
    with store.transaction() as s:
        collection = resource(s, data={
            "title": "Full", "question_ids": [],
            "items": [{"number": index, "question_id": "q-%s" % index,
                       "status": "detached", "batch_id": "old"}
                      for index in range(1, 1001)],
            "batches": [], "next_question_number": 1001,
        })
        collection_id = collection.id
    response = client.post("/auto-task/ai/question-workflows/intake", json={
        "collection_id": collection_id, "title": "Overflow", "text": "one more",
    })
    assert response.status_code == 409
    with store.transaction() as s:
        assert s.query(TaskJob).count() == 0


def test_question_intake_runs_solve_verify_chain_without_code_execution(client, monkeypatch):
    from auto_task.runtime import Context
    question = {"question_type": "short_answer", "language": "zh", "original": "FIFO?", "zh": "什么是 FIFO？",
                "options": [], "answer": "", "explanation": "", "tags": []}
    def ask_model(self, prompt, schema, images=()):
        if schema["title"] == "Moderation":
            return {"decision": "normal", "reason": ""}
        if schema["title"] == "Extraction":
            return {"questions": [question]}
        if schema["title"] == "Answer":
            return {"answer": "先进先出", "explanation": "队列先进入先离开", "tags": ["数据结构"]}
        return {"passed": True, "evidence": ["答案符合定义"], "corrected_answer": "", "explanation": "通过",
                "requires_execution": True, "verification_code": "print(1)", "language": "python",
                "tests": [{"input": "", "output": "1"}]}
    monkeypatch.setattr(Context, "ask", ask_model)
    response = client.post("/auto-task/ai/question-workflows/intake", json={
        "title": "Chain", "text": "FIFO", "batch_id": "chain-1", "auto_solve": True, "auto_verify": True})
    assert response.status_code == 200, response.text
    assert run_once()
    with store.transaction() as s:
        collection = s.query(AiResource).filter_by(kind="collection").one()
        assert len(json.loads(collection.data)["question_ids"]) == 1
        solve = s.query(TaskJob).filter_by(task_type="question.solve").one()
        assert solve.status == "pending"
    assert run_once()
    with store.transaction() as s:
        question_row = s.query(AiResource).filter_by(kind="question").one()
        solve = s.query(TaskJob).filter_by(task_type="question.solve").one()
        assert solve.status == "success"
        assert json.loads(question_row.data)["answer"] == "先进先出"
        verify = s.query(TaskJob).filter_by(task_type="question.verify").one()
        assert json.loads(verify.payload)["allow_execution"] is False
    assert run_once()
    with store.transaction() as s:
        verify = s.query(TaskJob).filter_by(task_type="question.verify").one()
        assert verify.status == "success"
        result = json.loads(verify.result)
        assert result["execution_blocked"] is True


def test_question_followup_and_revision_restore_are_versioned(client):
    client.app.dependency_overrides[api.cover_header] = lambda: {**USER, "roles": ["admin"]}
    created = client.post("/auto-task/ai/resources/create", json={"kind": "question", "data": {
        "question_type": "short_answer", "language": "zh", "original": "old", "zh": "旧题",
        "options": [], "answer": "A", "tags": []}}).json()["data"]
    updated = client.post("/auto-task/ai/resources/%s/update" % created["id"], json={
        "version": created["version"], "data": {"zh": "新题"}})
    assert updated.status_code == 200, updated.text
    followup = client.post("/auto-task/ai/question-workflows/questions/%s/followup" % created["id"], json={
        "operation_key": "solve-1", "task_type": "question.solve"})
    assert followup.status_code == 200, followup.text
    assert client.post("/auto-task/ai/question-workflows/questions/%s/followup" % created["id"], json={
        "operation_key": "chat-empty", "task_type": "question.chat"}).status_code == 400
    chat = client.post("/auto-task/ai/question-workflows/questions/%s/followup" % created["id"], json={
        "operation_key": "chat-1", "task_type": "question.chat", "message": "为什么？"})
    assert chat.status_code == 200, chat.text
    with store.transaction() as s:
        assert json.loads(s.query(TaskJob).get(chat.json()["data"]["taskId"]).payload)["message"] == "为什么？"
    history = client.get("/auto-task/ai/resources/%s/history" % created["id"]).json()["data"]
    assert history["totalNum"] >= 2
    assert client.post("/auto-task/ai/resources/%s/delete" % created["id"]).status_code == 200
    restored = client.post("/auto-task/ai/question-workflows/questions/%s/restore" % created["id"], json={
        "version": created["version"], "revision_id": next(row["id"] for row in history["rows"] if row["version"] == created["version"])})
    assert restored.status_code == 200, restored.text
    assert restored.json()["data"]["data"]["zh"] == "旧题"
    with store.transaction() as s:
        assert s.query(AiResource).get(created["id"]).deleted == 0


def test_question_revision_restore_replaces_file_links(client):
    client.app.dependency_overrides[api.cover_header] = lambda: {**USER, "roles": ["admin"]}
    created = client.post("/auto-task/ai/resources/create", json={"kind": "question", "data": {
        "question_type": "short_answer", "language": "zh", "original": "old", "zh": "旧题",
        "options": [], "answer": "A", "tags": []}}).json()["data"]
    with store.transaction() as s:
        s.add(AiFile(id="historic-file", owner_id="42", resource_id=created["id"],
                     filename="figure.png", media_type="image/png", path="historic.png",
                     sha256="a" * 64, size=10))
    with_file = client.post("/auto-task/ai/resources/%s/update" % created["id"], json={
        "version": created["version"], "data": {"zh": "带图版本", "file_ids": ["historic-file"]}}).json()["data"]
    without_file = client.post("/auto-task/ai/resources/%s/update" % created["id"], json={
        "version": with_file["version"], "data": {"zh": "无图版本", "file_ids": []}})
    assert without_file.status_code == 200, without_file.text
    history = client.get("/auto-task/ai/resources/%s/history" % created["id"]).json()["data"]
    revision = next(row for row in history["rows"] if row["version"] == with_file["version"])

    restored = client.post("/auto-task/ai/question-workflows/questions/%s/restore" % created["id"], json={
        "version": revision["version"], "revision_id": revision["id"]})
    assert restored.status_code == 200, restored.text
    assert restored.json()["data"]["data"]["file_ids"] == ["historic-file"]
    with store.transaction() as s:
        assert [(link.resource_id, link.file_id) for link in s.query(AiFileLink).all()] == [
            (created["id"], "historic-file")]


def test_question_intake_rejects_another_users_source_file(client):
    with store.transaction() as s:
        s.add(AiFile(id="foreign", owner_id="7", filename="private.pdf", media_type="application/pdf",
                     path="private.pdf", sha256="e" * 64, size=10))
    response = client.post("/auto-task/ai/question-workflows/intake", json={
        "title": "Foreign", "file_ids": ["foreign"]})
    assert response.status_code == 403


def test_question_tags_are_typed_and_bounded(client):
    invalid = client.post("/auto-task/ai/resources/create", json={"kind": "question", "data": {
        "question_type": "short_answer", "language": "zh", "original": "Q", "zh": "Q",
        "options": [], "answer": "", "tags": "not-a-list",
    }})
    assert invalid.status_code == 400
    too_many = client.post("/auto-task/ai/resources/create", json={"kind": "question", "data": {
        "question_type": "short_answer", "language": "zh", "original": "Q", "zh": "Q",
        "options": [], "answer": "", "tags": ["t%s" % index for index in range(7)],
    }})
    assert too_many.status_code == 400


def native_answer_fixture(client):
    client.app.dependency_overrides[api.cover_header] = lambda: ADMIN
    with store.transaction() as s:
        s.add(ProblemSet(psid=1, name="exam", type=1, groupInfo='[{"gid":1}]', config="{}", username="teacher", groupId=1, tag="test"))
        s.add(ProblemGroup(gid=1, name="subjective", type=1, problemInfo='[{"pid":1,"score":10}]', username="teacher"))
        s.add(ProblemSubjective(pid=1, type=1, config='{"judgeConfig":[{"name":"correct","score":10}]}', description="question", username="teacher"))
        s.add(ProblemSetAnswerSheet(asid=1, psid=1, username="student"))
        s.flush()
        s.add(ProblemSetAnswerSheetDetail(asd_id=1, asid=1, gid=1, pid=1, answer='["answer"]'))
    return {"psid": 1, "gid": 1, "pid": 1, "username": "student", "operation_key": "native"}


def test_native_grade_submission_freezes_source_and_rejects_stale_input(client):
    body = native_answer_fixture(client)
    first = client.post("/auto-task/ai/native/grading/create", json=body)
    assert first.status_code == 200, first.text
    result = first.json()["data"]
    assert client.post("/auto-task/ai/native/grading/create", json=body).json()["data"]["taskId"] == result["taskId"]
    assert client.post("/auto-task/ai/native/grading/create", json={**body, "gid": 2}).status_code == 403
    with store.transaction() as s:
        s.query(ProblemSetAnswerSheetDetail).get(1).answer = '["changed"]'
    assert run_once()
    with store.transaction() as s:
        assert s.query(TaskJob).get(result["taskId"]).status == "failed"
        assert not s.query(ProblemSetAnswerSheetDetail).get(1).judgeLog




def test_exam_publication_is_bound_to_reviewed_native_payload(client, monkeypatch):
    from server import auto_task_actions as actions
    client.app.dependency_overrides[api.cover_header] = lambda: ADMIN
    async def creator(*args):
        return None
    monkeypatch.setattr("server.auto_task._assert_group_creator", creator)
    monkeypatch.setattr(actions, "group_manager", lambda *args: None)
    with store.transaction() as s:
        resource(s, "7", "exam", {"mode": "TEST", "candidate": {"rooms": [], "candidate_count": 0}})
    native = {"name": "Dev test", "type": 1, "groupId": 10, "tag": "dev",
        "groupInfo": [{"gid": 1, "name": "group", "score": 100}], "tm_start": 1788610000000, "tm_end": 1788620000000,
        "config": {"useSameSE": 1, "showReport": 0, "showScoreInRunning": 0, "showProgramScoreInRunning": 0, "mergerSubjectiveGroup": 0}}
    path = "/auto-task/ai/exams/res-7"
    prepared = client.post(path + "/prepare", json={"version": 1, "problem_set": native})
    assert prepared.status_code == 200, prepared.text
    assert client.post(path + "/publish", json={"version": 2, "problem_set": native, "dry_run": False}).status_code == 409
    assert client.post("/auto-task/ai/resources/res-7/review", json={"version": 2, "reason": "checked payload"}).status_code == 200
    assert client.post(path + "/publish", json={"version": 3, "problem_set": {**native, "name": "changed"}, "dry_run": False}).status_code == 409
    published = client.post(path + "/publish", json={"version": 3, "problem_set": native, "dry_run": False})
    assert published.status_code == 200, published.text
    assert client.post(path + "/publish", json={"version": 3, "problem_set": native, "dry_run": False}).json()["data"]["replay"]


def test_pair_review_and_report_are_authorized_and_versioned(client):
    from auto_task.configuration import safe_path
    root = safe_path("workspaces/res-42")
    root.mkdir(parents=True)
    (root / "report.jplag").write_bytes(b"report-fixture")
    with store.transaction() as s:
        resource(s, "42", "plagiarism", {"candidate": {"report": "report.jplag", "pairs": [{"left": "one", "right": "two", "similarity": 0.9}]}})
        resource(s, "99", "plagiarism", {"candidate": {"report": "../../res-42/report.jplag"}})
    path = "/auto-task/ai/plagiarism/res-42"
    response = client.post(path + "/pairs/review", json={"version": 1, "left": "one", "right": "two", "decision": "CLEAR", "reason": "shared template"})
    assert response.json()["data"] == {"version": 2, "decision": "CLEAR", "native_grade_changed": False}
    assert client.get(path + "/report").content == b"report-fixture"
    assert client.get("/auto-task/ai/plagiarism/res-99/report").status_code == 403
    with store.transaction() as s:
        s.query(AiResource).get("res-42").data = '{"candidate":{"report":"../res-99/secret"}}'
    assert client.get(path + "/report").status_code == 404


@pytest.mark.parametrize("group_type", [0, 1])
def test_true_false_publish_uses_only_native_objective_groups(client, monkeypatch, group_type):
    from server import auto_task_actions as actions
    monkeypatch.setattr(actions, "group_manager", lambda *args: None)
    with store.transaction() as session:
        session.add(ProblemGroup(gid=1, name="target", type=group_type, problemInfo="[]", username="student"))
        resource(session, "42", "question", {"question_type": "true_false", "original": "1=1", "zh": "1=1",
                    "options": [{"label": "A", "zh": "正确"}, {"label": "B", "zh": "错误"}]})
    body = {"version": 1, "group_id": 1, "score": 5, "correct_labels": ["A"], "dry_run": True}
    path = "/auto-task/ai/questions/res-42/publish"
    preview = client.post(path, json=body)
    assert preview.status_code == (200 if group_type == 0 else 400), preview.text
    result = client.post(path, json={**body, "dry_run": False})
    assert result.status_code == (200 if group_type == 0 else 400), result.text
    with store.transaction() as session:
        assert session.query(ProblemSubjective).count() == 0
        if group_type == 0:
            assert preview.json()["data"]["native_type"] == "objective"
            item = session.query(ProblemObjective).one()
            assert item.type == 0 and json.loads(item.answer) == ["A"]
        else:
            assert session.query(ProblemObjective).count() == 0


def test_true_false_publish_requires_exactly_one_answer(client, monkeypatch):
    from server import auto_task_actions as actions
    monkeypatch.setattr(actions, "group_manager", lambda *args: None)
    with store.transaction() as session:
        session.add(ProblemGroup(gid=1, name="choices", type=0, problemInfo="[]", username="student"))
        resource(session, "42", "question", {"question_type": "true_false", "zh": "Q",
                 "options": [{"label": "A", "zh": "yes"}, {"label": "B", "zh": "no"}]})
    response = client.post("/auto-task/ai/questions/res-42/publish", json={
        "version": 1, "group_id": 1, "score": 5, "correct_labels": ["A", "B"], "dry_run": False})
    assert response.status_code == 400
    with store.transaction() as session:
        assert session.query(ProblemObjective).count() == 0


def test_complete_clears_previous_blocking_error(database):
    job_id = enqueue()
    active = store.claim()
    with store.transaction() as s:
        s.query(TaskJob).filter_by(id=job_id).one().error = "AI usage reconciliation required"
    store.complete(job_id, active["lease_token"], {"ok": True})
    with store.transaction() as s:
        job = s.query(TaskJob).filter_by(id=job_id).one()
        assert job.status == "success"
        assert job.error is None


def test_release_provider_allows_only_configured_private_http_origin(monkeypatch):
    from pydantic import ValidationError

    monkeypatch.setattr("const.MOD", "release")
    monkeypatch.setenv("AUTOLAB_AI_TRUSTED_HTTP_PROVIDER_ORIGINS", "http://10.102.34.150:18080")

    config = ProviderConfig(
        provider="codex",
        base_url="http://10.102.34.150:18080/v1",
        model="deepseek-v4.1-flash",
    )
    assert config.base_url == "http://10.102.34.150:18080/v1"

    with pytest.raises(ValidationError):
        ProviderConfig(base_url="http://10.102.34.151:18080/v1", model="test")
    with pytest.raises(ValidationError):
        ProviderConfig(base_url="http://10.102.34.150:18080/private", model="test")


def test_responses_model_call_sends_stateless_images_and_settles_usage(database, monkeypatch, tmp_path):
    from auto_task.runtime import Context
    config = ProviderConfig(**{**model_config().dict(), "provider": "responses"})
    with store.transaction() as s:
        save_provider(s, config, "test-only-key")
    job_id = enqueue()
    job = store.claim()
    image = tmp_path / "sample.png"
    image.write_bytes(b"synthetic-image")
    captured = {}

    class Response:
        ok = True

        def iter_lines(self, decode_unicode=True):
            events = [
                {"type": "response.reasoning_summary_text.delta", "delta": "ignored"},
                {"type": "response.output_text.delta", "delta": '{"answer":'},
                {"type": "response.output_text.delta", "delta": '"4"}'},
                {"type": "response.completed", "response": {"status": "completed", "usage": {
                    "input_tokens": 100, "output_tokens": 10,
                    "input_tokens_details": {"cached_tokens": 20}}}},
            ]
            yield from ("data: " + json.dumps(event) for event in events)
            yield "data: [DONE]"

        def close(self):
            pass

    def post(url, **kwargs):
        captured.update(url=url, **kwargs)
        return Response()

    monkeypatch.setattr("auto_task.runtime.requests.post", post)
    schema = {"type": "object", "properties": {"answer": {"type": "string"}}, "required": ["answer"]}
    assert Context(job).ask("2+2?", schema, [image]) == {"answer": "4"}
    assert captured["url"] == "https://example.invalid/v1/responses"
    payload = captured["json"]
    assert payload["store"] is False and payload["stream"] is True
    assert "previous_response_id" not in payload
    visible_prompt = payload["input"][0]["content"][0]["text"]
    assert store.dumps(schema) in visible_prompt and "2+2?" in visible_prompt
    assert "Return exactly one JSON object" in visible_prompt
    assert payload["input"][0]["content"][1]["type"] == "input_image"
    assert payload["text"]["format"]["type"] == "json_schema"
    assert captured["allow_redirects"] is False
    with store.transaction() as s:
        run = s.query(AiRun).filter_by(task_id=job_id).one()
        assert run.status == "success"
        assert json.loads(run.usage)["cached_input_tokens"] == 20
        assert run.cost_usd == Decimal("0.00011")


@pytest.mark.parametrize("status,content", [
    ("incomplete", [{"type": "output_text", "text": '{"answer":"4"}'}]),
    ("failed", []),
    ("completed", [{"type": "refusal", "refusal": "No"}]),
    ("completed", []),
])
def test_responses_rejects_incomplete_or_refused_outputs(monkeypatch, status, content):
    from auto_task.runtime import Context
    ctx = object.__new__(Context)
    config = ProviderConfig(**{**model_config().dict(), "provider": "responses"})

    class Response:
        ok = True

        def iter_lines(self, decode_unicode=True):
            usage = {"input_tokens": 10, "output_tokens": 1}
            output = [{"type": "message", "content": content}] if content else []
            if status == "completed":
                event = {"type": "response.completed", "response": {"status": status, "output": output, "usage": usage}}
                if content and content[0].get("type") == "output_text":
                    event = {"type": "response.output_text.delta", "delta": content[0]["text"]}
                    yield "data: " + json.dumps(event)
                    event = {"type": "response.completed", "response": {"status": status, "output": output, "usage": usage}}
            else:
                event = {"type": "response." + status, "response": {"status": status, "output": output, "usage": usage}}
            yield "data: " + json.dumps(event)

        def close(self):
            pass

    monkeypatch.setattr("auto_task.runtime.requests.post", lambda *_, **kwargs: Response())
    with pytest.raises(ValueError):
        ctx._responses(config, "test-only-key", "test", {}, [])
    assert ctx.observed_usage == {"input_tokens": 10, "output_tokens": 1}


def test_responses_rejects_prose_while_retaining_observed_usage(monkeypatch):
    from auto_task.runtime import Context
    ctx = object.__new__(Context)
    config = ProviderConfig(**{**model_config().dict(), "provider": "responses"})

    class Response:
        ok = True

        def iter_lines(self, decode_unicode=True):
            yield "data: " + json.dumps({"type": "response.output_text.delta", "delta": "**Classification: Normal**"})
            yield "data: " + json.dumps({"type": "response.completed", "response": {
                "status": "completed", "usage": {"input_tokens": 10, "output_tokens": 1}}})

        def close(self):
            pass

    monkeypatch.setattr("auto_task.runtime.requests.post", lambda *_, **kwargs: Response())
    with pytest.raises(ValueError, match="did not contain valid JSON"):
        ctx._responses(config, "test-only-key", "test", {}, [])
    assert ctx.observed_usage == {"input_tokens": 10, "output_tokens": 1}


def test_responses_accepts_schema_json_when_compatible_gateway_omits_terminal_event(monkeypatch):
    from auto_task.runtime import Context
    ctx = object.__new__(Context)
    config = ProviderConfig(**{**model_config().dict(), "provider": "responses"})

    class Response:
        ok = True

        def iter_lines(self, decode_unicode=True):
            yield "data: " + json.dumps({"type": "response.output_text.delta", "delta": '{"answer":"4"}'})

        def close(self):
            pass

    monkeypatch.setattr("auto_task.runtime.requests.post", lambda *_, **kwargs: Response())
    result, usage = ctx._responses(config, "test-only-key", "test", {
        "type": "object", "properties": {"answer": {"type": "string"}}, "required": ["answer"],
    }, [])
    assert result == {"answer": "4"}
    assert usage == {} and ctx.observed_usage == {}
    assert ctx.stream_diagnostics["accepted_eof_json"] is True


@pytest.mark.parametrize("terminal", [False, True])
def test_responses_decodes_chinese_sse_as_utf8(monkeypatch, terminal):
    from auto_task.runtime import Context
    ctx = object.__new__(Context)
    config = ProviderConfig(**{**model_config().dict(), "provider": "responses"})
    expected = {"name": "算法正确性", "summary": "答案正确"}

    class Response:
        ok = True
        encoding = "ISO-8859-1"

        def iter_lines(self, decode_unicode=True):
            events = [{"type": "response.output_text.delta", "delta": json.dumps(expected, ensure_ascii=False)}]
            if terminal:
                events.append({"type": "response.completed", "response": {
                    "status": "completed", "usage": {"input_tokens": 10, "output_tokens": 2}}})
            for event in events:
                yield ("data: " + json.dumps(event, ensure_ascii=False)).encode("utf-8").decode(self.encoding)

        def close(self):
            pass

    monkeypatch.setattr("auto_task.runtime.requests.post", lambda *_, **__: Response())
    result, _ = ctx._responses(config, "test-only-key", "test", {}, [])
    assert result == expected


def test_responses_5xx_does_not_fall_back_to_chat_completions(database, monkeypatch):
    from auto_task.runtime import Context
    config = ProviderConfig(**{**model_config().dict(), "provider": "responses"})
    with store.transaction() as s:
        save_provider(s, config, "test-only-key")
    job_id = enqueue()
    job = store.claim()
    captured = {}

    class Response:
        ok, status_code = False, 502

        def close(self):
            pass

    def post(url, **kwargs):
        captured.update(url=url, **kwargs)
        return Response()

    monkeypatch.setattr("auto_task.runtime.requests.post", post)
    schema = {"type": "object", "properties": {"answer": {"type": "string"}}, "required": ["answer"]}
    with pytest.raises(RuntimeError, match="HTTP 502 \(Responses\)"):
        Context(job).ask("2+2?", schema)
    assert captured["url"] == "https://example.invalid/v1/responses"
    assert captured["stream"] is True
    with store.transaction() as s:
        run = s.query(AiRun).filter_by(task_id=job_id).one()
        assert run.status == "unsettled"
        failure_logs = [json.loads(row.content) for row in s.query(AutoTaskRunLog).filter_by(task_id=job_id, tag="log")]
        failure = next(item for item in failure_logs if item.get("event") == "model_request_failed")
        assert failure["error_type"] == "ProviderHTTPError"
        assert failure["stream"]["event_types"] == []


def test_native_grade_worker_writes_back_and_teacher_can_override(client, monkeypatch):
    import asyncio
    from types import SimpleNamespace
    from unittest.mock import AsyncMock, MagicMock
    from model.answer_sheet import answerSheetModel
    from auto_task.runtime import Context
    import auto_task.subjective_review as legacy
    body = native_answer_fixture(client)
    result = client.post("/auto-task/ai/native/grading/create", json=body).json()["data"]
    service = SimpleNamespace(loop=None, payload=None, answer_model=MagicMock(), subject_model=MagicMock(),
        _run_async=lambda x: x, _get_answer_detail_dict=lambda: {},
        _get_student_answer_text=lambda *args: "answer", _get_subject_markdown=lambda *args: "question",
        _build_programming_sections=lambda *args: "")
    service.answer_model.get_gid_pid2indexDict_cache.return_value = ({}, {})
    monkeypatch.setattr(legacy, "_SubjectiveReviewService", lambda *args: service)
    monkeypatch.setattr(legacy, "_ensure_fastapi_cache", lambda *args: None)
    monkeypatch.setattr("sduojApi.getUserId", lambda *args: "42")
    monkeypatch.setattr(Context, "ask", lambda self, prompt, schema, images=(): {"decision": "normal", "reason": ""} if schema["title"] == "Moderation" else {
        "criteria": [{"name": "correct", "max_score": 10, "score": 8, "evidence": "answer"}],
        "summary": "checked", "needs_review": False, "confidence": 1})
    assert run_once()
    with store.transaction() as s:
        job = s.query(TaskJob).get(result["taskId"])
        assert job.status == "success", job.error
        detail = s.query(ProblemSetAnswerSheetDetail).get(1)
        assert json.loads(detail.judgeLog)[0]["jScore"] == 8
        assert detail.judgeComment == "checked"
        assert detail.judgeLock_username == "<AUTO>"
        assert json.loads(job.result)["written_back"] is True
        assert json.loads(job.result)["candidate"] is False
        assert s.query(AiResource).get(result["resource_id"]).state == "adopted"
        assert s.query(AiResource).get(result["resource_id"]).version == 2
    adopted = client.post("/auto-task/ai/grading/" + result["resource_id"] + "/adopt", json={"version": 2, "reason": "Teacher checked evidence"})
    assert adopted.status_code == 200, adopted.text
    assert adopted.json()["data"]["replay"] is True

    # The existing manual grading model can still edit the automatic score.
    manual = object.__new__(answerSheetModel)
    manual.session = store.factory()()
    manual.get_gid_pid_by_psid_gi_pi_cache = AsyncMock(return_value=(1, 1))
    try:
        manual.update_detail_by_asd_id(1, {"judgeLock_username": "teacher"})
        asyncio.run(manual.update_judgeLog_by_psid_gi_pi_username(
            1, 0, 0, "student", [{"name": "correct", "score": 10, "jScore": 9}], 0, "teacher corrected"))
    finally:
        manual.session.close()
    with store.transaction() as s:
        detail = s.query(ProblemSetAnswerSheetDetail).get(1)
        assert json.loads(detail.judgeLog)[0]["jScore"] == 9
        assert detail.judgeComment == "teacher corrected"
        from auto_task.legacy_bridge import writeback_subjective_result
        completed = s.query(TaskJob).get(result["taskId"])
        writeback_subjective_result(s, completed, json.loads(completed.result))
        assert json.loads(detail.judgeLog)[0]["jScore"] == 9

    # A later automatic run preserves a grade already claimed by a teacher.
    monkeypatch.setattr(Context, "ask", lambda *_args, **_kwargs: pytest.fail("manual grade must skip model calls"))
    retry = client.post("/auto-task/ai/native/grading/create", json={**body, "operation_key": "after-manual"})
    assert retry.status_code == 200, retry.text
    assert run_once()
    with store.transaction() as s:
        job = s.query(TaskJob).get(retry.json()["data"]["taskId"])
        assert job.status == "success", job.error
        assert json.loads(job.result)["skipped"] == "manual_grade_exists"
        assert json.loads(s.query(ProblemSetAnswerSheetDetail).get(1).judgeLog)[0]["jScore"] == 9


@pytest.mark.parametrize("change", ["answer", "grade", "rubric"])
def test_native_auto_writeback_preserves_edits_during_model_execution(client, change):
    from auto_task.legacy_bridge import answer_snapshot, writeback_subjective_result
    body = native_answer_fixture(client)
    created = client.post("/auto-task/ai/native/grading/create", json=body).json()["data"]
    with store.transaction() as s:
        _, source = answer_snapshot(s, 1, 1, 1, "student")
        value = {"source": source, "judgeLog": [{"name": "correct", "score": 10, "jScore": 8}],
                 "judgeComment": "checked", "evaluation": {"criteria": [{"name": "correct", "max_score": 10,
                    "score": 8, "evidence": "answer"}], "summary": "checked", "needs_review": False, "confidence": 1}}
        detail = s.query(ProblemSetAnswerSheetDetail).get(1)
        if change == "answer":
            detail.answer = '["new answer"]'
        elif change == "grade":
            detail.judgeLog = '[{"name":"correct","score":10,"jScore":9}]'
            detail.judgeComment, detail.judgeLock_username = "teacher correction", "teacher"
        else:
            s.query(ProblemSubjective).get(1).config = '{"judgeConfig":[{"name":"changed","score":10}]}'
    with pytest.raises(ValueError, match="changed during model execution"):
        with store.transaction() as s:
            writeback_subjective_result(s, s.query(TaskJob).get(created["taskId"]), value)
    with store.transaction() as s:
        detail = s.query(ProblemSetAnswerSheetDetail).get(1)
        assert detail.judgeLock_username != "<AUTO>"
        if change == "grade":
            assert json.loads(detail.judgeLog)[0]["jScore"] == 9
        else:
            assert not detail.judgeLog
